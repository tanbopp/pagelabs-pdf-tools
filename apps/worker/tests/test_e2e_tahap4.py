import requests
import json
import time
from pathlib import Path
from PIL import Image
import docx
import openpyxl
from pptx import Presentation
from pptx.util import Inches, Pt

API_BASE = "http://localhost:3001"

def create_sample_docx(path: Path):
    doc = docx.Document()
    doc.add_heading("PageLabs Tahap 4 Word Test", level=1)
    doc.add_paragraph("This is a test paragraph inside Word DOCX converting to PDF.")
    doc.save(str(path))

def create_sample_xlsx(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "Item"
    ws["B1"] = "Quantity"
    ws["A2"] = "Widget A"
    ws["B2"] = 100
    ws["A3"] = "Widget B"
    ws["B3"] = 250
    wb.save(str(path))

def create_sample_pptx(path: Path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "PageLabs PPTX Conversion"
    subtitle.text = "Automated E2E Test"
    prs.save(str(path))

def create_sample_jpg(path: Path):
    img = Image.new("RGB", (400, 300), color=(73, 109, 137))
    img.save(str(path), format="JPEG")

def create_sample_html(path: Path):
    html_content = """<!DOCTYPE html>
<html>
<head>
  <meta charset="utf-8">
  <title>E2E HTML Document</title>
  <style>
    body { font-family: sans-serif; padding: 24px; background: #fafafa; }
    h1 { color: #2563eb; }
    p { color: #374151; font-size: 14px; }
  </style>
</head>
<body>
  <h1>PageLabs HTML to PDF</h1>
  <p>Rendered with headless Chromium for pixel perfection.</p>
</body>
</html>"""
    path.write_text(html_content, encoding="utf-8")

def upload_file(file_path: Path, content_type: str):
    size = file_path.stat().st_size
    sign_res = requests.post(f"{API_BASE}/api/v1/uploads/sign", json={
        "filename": file_path.name,
        "size": size,
        "contentType": content_type
    })
    sign_res.raise_for_status()
    signed = sign_res.json()
    upload_url = signed["uploadUrl"]
    file_url = signed["fileUrl"]

    with open(file_path, "rb") as f:
        put_res = requests.put(upload_url, data=f, headers={"Content-Type": content_type})
        put_res.raise_for_status()

    return {"url": file_url, "name": file_path.name, "size": size}

def poll_job(job_id: str, max_wait=40):
    start = time.time()
    while time.time() - start < max_wait:
        res = requests.get(f"{API_BASE}/api/v1/jobs/{job_id}")
        res.raise_for_status()
        data = res.json()
        if data["status"] in ("completed", "failed"):
            return data
        time.sleep(1)
    raise TimeoutError(f"Job {job_id} timed out")

def run():
    tmp = Path("apps/api/storage/uploads/e2e_tahap4_tmp")
    tmp.mkdir(parents=True, exist_ok=True)

    docx_path = tmp / "sample.docx"
    xlsx_path = tmp / "sample.xlsx"
    pptx_path = tmp / "sample.pptx"
    jpg_path = tmp / "sample.jpg"
    html_path = tmp / "sample.html"

    create_sample_docx(docx_path)
    create_sample_xlsx(xlsx_path)
    create_sample_pptx(pptx_path)
    create_sample_jpg(jpg_path)
    create_sample_html(html_path)

    # 1. Word to PDF
    print("Testing word-to-pdf...")
    up_docx = upload_file(docx_path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    j_res = requests.post(f"{API_BASE}/api/v1/jobs", json={
        "tool": "word-to-pdf",
        "files": [{"url": up_docx["url"], "name": up_docx["name"], "size": up_docx["size"]}],
        "options": {"outputFilename": "word_converted.pdf"}
    })
    j_res.raise_for_status()
    status_data = poll_job(j_res.json()["jobId"])
    assert status_data["status"] == "completed", f"Word to PDF failed: {status_data}"
    print("word-to-pdf PASSED! Result:", status_data["result_url"])

    # 2. Excel to PDF
    print("Testing excel-to-pdf...")
    up_xlsx = upload_file(xlsx_path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    j_res = requests.post(f"{API_BASE}/api/v1/jobs", json={
        "tool": "excel-to-pdf",
        "files": [{"url": up_xlsx["url"], "name": up_xlsx["name"], "size": up_xlsx["size"]}],
        "options": {"fitToPage": True, "outputFilename": "excel_converted.pdf"}
    })
    j_res.raise_for_status()
    status_data = poll_job(j_res.json()["jobId"])
    assert status_data["status"] == "completed", f"Excel to PDF failed: {status_data}"
    print("excel-to-pdf PASSED! Result:", status_data["result_url"])

    # 3. PowerPoint to PDF
    print("Testing powerpoint-to-pdf...")
    up_pptx = upload_file(pptx_path, "application/vnd.openxmlformats-officedocument.presentationml.presentation")
    j_res = requests.post(f"{API_BASE}/api/v1/jobs", json={
        "tool": "powerpoint-to-pdf",
        "files": [{"url": up_pptx["url"], "name": up_pptx["name"], "size": up_pptx["size"]}],
        "options": {"outputFilename": "ppt_converted.pdf"}
    })
    j_res.raise_for_status()
    status_data = poll_job(j_res.json()["jobId"])
    assert status_data["status"] == "completed", f"PowerPoint to PDF failed: {status_data}"
    print("powerpoint-to-pdf PASSED! Result:", status_data["result_url"])

    # 4. JPG to PDF
    print("Testing jpg-to-pdf...")
    up_jpg = upload_file(jpg_path, "image/jpeg")
    j_res = requests.post(f"{API_BASE}/api/v1/jobs", json={
        "tool": "jpg-to-pdf",
        "files": [{"url": up_jpg["url"], "name": up_jpg["name"], "size": up_jpg["size"]}],
        "options": {"orientation": "portrait", "pageSize": "a4", "margin": "small", "outputFilename": "jpg_converted.pdf"}
    })
    j_res.raise_for_status()
    status_data = poll_job(j_res.json()["jobId"])
    assert status_data["status"] == "completed", f"JPG to PDF failed: {status_data}"
    print("jpg-to-pdf PASSED! Result:", status_data["result_url"])

    # 5. HTML to PDF
    print("Testing html-to-pdf...")
    up_html = upload_file(html_path, "text/html")
    j_res = requests.post(f"{API_BASE}/api/v1/jobs", json={
        "tool": "html-to-pdf",
        "files": [{"url": up_html["url"], "name": up_html["name"], "size": up_html["size"]}],
        "options": {"pageSize": "a4", "outputFilename": "html_converted.pdf"}
    })
    j_res.raise_for_status()
    status_data = poll_job(j_res.json()["jobId"])
    assert status_data["status"] == "completed", f"HTML to PDF failed: {status_data}"
    print("html-to-pdf PASSED! Result:", status_data["result_url"])

    print("\nALL 5 TAHAP 4 E2E CONVERSION JOBS PASSED 100% GREEN!")

if __name__ == "__main__":
    run()
