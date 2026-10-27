import pytest
from pathlib import Path
from PIL import Image
import docx
import openpyxl
import pptx
import pikepdf

from recipes.word_to_pdf import word_to_pdf_recipe
from recipes.excel_to_pdf import excel_to_pdf_recipe
from recipes.ppt_to_pdf import ppt_to_pdf_recipe
from recipes.jpg_to_pdf import jpg_to_pdf_recipe
from recipes.html_to_pdf import html_to_pdf_recipe


@pytest.fixture
def sample_docx(tmp_path) -> Path:
    p = tmp_path / "test.docx"
    doc = docx.Document()
    doc.add_heading("Hello World Docx", level=1)
    doc.add_paragraph("This is a test paragraph for docx conversion.")
    doc.save(str(p))
    return p


@pytest.fixture
def sample_xlsx(tmp_path) -> Path:
    p = tmp_path / "test.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["ID", "Name", "Score"])
    ws.append([1, "Alice", 95])
    ws.append([2, "Bob", 88])
    wb.save(str(p))
    return p


@pytest.fixture
def sample_pptx(tmp_path) -> Path:
    p = tmp_path / "test.pptx"
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Hello PPTX"
    slide.placeholders[1].text = "Subtitle slide test"
    prs.save(str(p))
    return p


@pytest.fixture
def sample_jpgs(tmp_path) -> list[Path]:
    paths = []
    for color, name in [((255, 0, 0), "red.jpg"), ((0, 255, 0), "green.jpg")]:
        img = Image.new("RGB", (200, 200), color=color)
        p = tmp_path / name
        img.save(p, format="JPEG")
        paths.append(p)
    return paths


@pytest.fixture
def sample_html(tmp_path) -> Path:
    p = tmp_path / "test.html"
    p.write_text("<!DOCTYPE html><html><body><h1>PageLabs HTML Test</h1><p>Rendered via Chromium</p></body></html>", encoding="utf-8")
    return p


def test_word_to_pdf(sample_docx, tmp_path):
    out = tmp_path / "doc.pdf"
    res = word_to_pdf_recipe(sample_docx, out)
    assert res.exists()
    assert res.stat().st_size > 0
    with pikepdf.open(res) as check:
        assert len(check.pages) >= 1


def test_excel_to_pdf(sample_xlsx, tmp_path):
    out = tmp_path / "sheet.pdf"
    res = excel_to_pdf_recipe(sample_xlsx, out, {"fitToPage": True})
    assert res.exists()
    assert res.stat().st_size > 0
    with pikepdf.open(res) as check:
        assert len(check.pages) >= 1


def test_ppt_to_pdf(sample_pptx, tmp_path):
    out = tmp_path / "slides.pdf"
    res = ppt_to_pdf_recipe(sample_pptx, out)
    assert res.exists()
    assert res.stat().st_size > 0
    with pikepdf.open(res) as check:
        assert len(check.pages) >= 1


def test_jpg_to_pdf(sample_jpgs, tmp_path):
    out = tmp_path / "gallery.pdf"
    res = jpg_to_pdf_recipe(sample_jpgs, out, {"orientation": "portrait", "pageSize": "a4", "margin": "small"})
    assert res.exists()
    assert res.stat().st_size > 0
    with pikepdf.open(res) as check:
        assert len(check.pages) == 2


def test_html_to_pdf(sample_html, tmp_path):
    out = tmp_path / "page.pdf"
    res = html_to_pdf_recipe(sample_html, out, {"pageSize": "a4"})
    assert res.exists()
    assert res.stat().st_size > 0
    with pikepdf.open(res) as check:
        assert len(check.pages) >= 1
