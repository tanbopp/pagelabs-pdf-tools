import shutil
import subprocess
from pathlib import Path
import pptx
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet


def find_soffice_binary() -> str | None:
    for cmd in ["soffice", "libreoffice"]:
        if shutil.which(cmd):
            return cmd
    return None


def convert_pptx_fallback(input_path: Path, output_path: Path) -> None:
    """Fallback converter using python-pptx and reportlab."""
    try:
        prs = pptx.Presentation(str(input_path))
    except Exception as e:
        raise ValueError("Failed to convert document: unreadable or corrupt file") from e

    pdf = SimpleDocTemplate(
        str(output_path),
        pagesize=landscape(letter),
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40,
    )
    styles = getSampleStyleSheet()
    story = []

    for i, slide in enumerate(prs.slides, start=1):
        story.append(Paragraph(f"<b>Slide {i}</b>", styles["Heading1"]))
        story.append(Spacer(1, 10))

        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    t = paragraph.text.strip()
                    if t:
                        slide_texts.append(t)

        if slide_texts:
            for text in slide_texts:
                safe_text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(safe_text, styles["Normal"]))
                story.append(Spacer(1, 4))
        else:
            story.append(Paragraph("[Slide contains graphics or empty text]", styles["Italic"]))

        story.append(Spacer(1, 24))

    if not story:
        story.append(Paragraph("[Empty Presentation]", styles["Normal"]))

    pdf.build(story)


def ppt_to_pdf_recipe(
    input_path: Path,
    output_path: Path,
    options: dict = None
) -> Path:
    """
    Convert PowerPoint (.pptx, .ppt) to PDF using LibreOffice Headless.
    Fallback: python-pptx & reportlab for .pptx.
    """
    soffice = find_soffice_binary()
    converted_via_soffice = False

    if soffice:
        cmd = [
            soffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_path.parent),
            str(input_path),
        ]
        try:
            res = subprocess.run(cmd, capture_output=True, timeout=120)
            expected_out = output_path.parent / f"{input_path.stem}.pdf"
            if res.returncode == 0 and expected_out.exists() and expected_out.stat().st_size > 0:
                if expected_out != output_path:
                    shutil.move(expected_out, output_path)
                converted_via_soffice = True
        except Exception:
            converted_via_soffice = False

    if not converted_via_soffice:
        ext = input_path.suffix.lower()
        if ext == ".pptx":
            convert_pptx_fallback(input_path, output_path)
        else:
            raise ValueError("Failed to convert document: LibreOffice required for legacy .ppt conversion")

    return output_path
